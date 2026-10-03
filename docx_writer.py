"""
docx_writer.py — Universal Enterprise DOCX Generation & Template Patching Engine.

Dual-Mode Architecture:
1. Generative Mode (JSON Snapshot / AST → DOCX):
   - High-fidelity paragraph & text run styling (font, size, color, spacing, indentation).
   - Real image embedding from local path, assets directory, or Base64 data with printable
     margin constraint check (ERR_DOCX_005) and graceful visual fallback cards.
   - Complex table generation supporting merged cells (colspan & rowspan) via native OpenXML
     grid mapping, vertical centering (w:vAlign="center"), cantSplit, tblHeader, and
     the Last Paragraph Rule (ERR_DOCX_001).
   - Multi-section layout with dynamic orientation switching (Portrait ↔ Landscape),
     custom margins, header/footer configuration, and dynamic page numbering.
2. In-Place Template Patching Mode (patch_docx_template):
   - Preserves 100% of existing document layout, Cover Page, shapes, and DrawingML.
   - Deep text placeholder replacement ({{KEY}}) across body, tables, headers, and footers
     with run-merging to prevent split-run token breakage.
   - Block anchor replacement ({{BLOCK:id}}) replacing placeholder paragraphs with
     dynamically constructed tables, images, or formatted paragraph sequences.
3. Surgical Content Injection (inject_content_into_docx, inject_diagram_into_docx):
   - Keeps backward compatibility for inserting diagrams and content after specific headings.
"""

from __future__ import annotations

import base64
import copy
import io
import json
import os
import re
import shutil
from typing import Any, Dict, List, Optional, Tuple, Union

from docx import Document
from docx.document import Document as DocumentType
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml

from PIL import Image as PILImage

try:
    from .smart_post_processor import calculate_heuristic_column_widths
except (ImportError, ValueError):
    from smart_post_processor import calculate_heuristic_column_widths

__all__ = [
    "write_docx",
    "write_docx_from_json_file",
    "patch_docx_template",
    "inject_content_into_docx",
    "inject_diagram_into_docx",
]

# Maximum sane Pt value Word accepts (≈ 1584pt = 22 inches)
_MAX_PT = 1584.0


def _safe_pt(value: float | None) -> Pt | None:
    """Return Pt(value) only if value is within Word's legal range."""
    if value is None:
        return None
    try:
        val_float = float(value)
    except (ValueError, TypeError):
        return None
    if abs(val_float) > _MAX_PT:
        return None
    return Pt(val_float)


# ---------------------------------------------------------------------------
# Formatting Helpers
# ---------------------------------------------------------------------------

_ALIGN_MAP: dict[str, WD_ALIGN_PARAGRAPH] = {
    "LEFT":       WD_ALIGN_PARAGRAPH.LEFT,
    "CENTER":     WD_ALIGN_PARAGRAPH.CENTER,
    "RIGHT":      WD_ALIGN_PARAGRAPH.RIGHT,
    "JUSTIFY":    WD_ALIGN_PARAGRAPH.JUSTIFY,
    "DISTRIBUTE": WD_ALIGN_PARAGRAPH.DISTRIBUTE,
}


def _hex_to_rgb(hex_color: str | None) -> RGBColor | None:
    if not hex_color:
        return None
    h = hex_color.lstrip("#").strip()
    if len(h) == 8:
        h = h[2:8]  # Strip ARGB leading alpha if present
    if len(h) != 6:
        return None
    try:
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    except ValueError:
        return None


def _apply_run_format(run, run_data: dict[str, Any]) -> None:
    """Apply formatting attributes from a run dict onto a python-docx Run."""
    if run_data.get("bold") is not None:
        run.bold = bool(run_data["bold"])
    if run_data.get("italic") is not None:
        run.italic = bool(run_data["italic"])
    if run_data.get("underline") is not None:
        run.underline = bool(run_data["underline"])
    if run_data.get("strike"):
        run.font.strike = True
    if run_data.get("font_name"):
        run.font.name = str(run_data["font_name"])
    size_pt = _safe_pt(run_data.get("font_size"))
    if size_pt is not None:
        run.font.size = size_pt
    color = _hex_to_rgb(run_data.get("font_color"))
    if color:
        run.font.color.rgb = color


def _set_keep_next(para) -> None:
    """Ensure paragraph has <w:keepNext/> so Word does not break page after it."""
    pPr = para._element.get_or_add_pPr()
    if pPr.find(qn("w:keepNext")) is None:
        pPr.append(OxmlElement("w:keepNext"))


def _apply_paragraph_format(para, para_data: dict[str, Any]) -> None:
    """Apply paragraph-level spacing/indent/alignment from a para dict."""
    fmt = para.paragraph_format
    alignment = _ALIGN_MAP.get(str(para_data.get("alignment", "LEFT")).upper(), WD_ALIGN_PARAGRAPH.LEFT)
    para.alignment = alignment

    sb = _safe_pt(para_data.get("space_before_pt"))
    if sb is not None:
        fmt.space_before = sb
    sa = _safe_pt(para_data.get("space_after_pt"))
    if sa is not None:
        fmt.space_after = sa
    li = _safe_pt(para_data.get("left_indent_pt"))
    if li is not None:
        fmt.left_indent = li
    fli = _safe_pt(para_data.get("first_line_indent_pt"))
    if fli is not None:
        fmt.first_line_indent = fli
    if para_data.get("line_spacing") is not None:
        ls = para_data["line_spacing"]
        if isinstance(ls, (int, float)) and ls > 4:
            safe = _safe_pt(ls)
            if safe is not None:
                fmt.line_spacing = safe
        else:
            fmt.line_spacing = ls

    if para_data.get("keep_next"):
        _set_keep_next(para)


def _resolve_style(doc: DocumentType, style_name: str):
    """
    Return the style object from doc matching style_name.
    Falls back to 'Normal' if not found.
    """
    try:
        return doc.styles[style_name]
    except KeyError:
        lower = style_name.lower()
        for s in doc.styles:
            if s.name.lower() == lower:
                return s
        return doc.styles["Normal"]


# ---------------------------------------------------------------------------
# OpenXML Table Formatting Helpers
# ---------------------------------------------------------------------------

def _set_cell_shading(tc_elem, hex_color: str | None) -> None:
    """Apply XML shading <w:shd> to a table cell."""
    if not hex_color:
        return
    h = hex_color.lstrip("#").upper()
    if len(h) == 8:
        h = h[2:8]
    if len(h) != 6:
        return
    tcPr = tc_elem.get_or_add_tcPr()
    for existing in tcPr.findall(qn("w:shd")):
        tcPr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), h)
    tcPr.append(shd)


def _set_cell_margins(tc_elem, top: int = 120, bottom: int = 120, left: int = 150, right: int = 150) -> None:
    """Set cell padding/margins via <w:tcMar> in dxa (1 pt = 20 dxa)."""
    tcPr = tc_elem.get_or_add_tcPr()
    for existing in tcPr.findall(qn("w:tcMar")):
        tcPr.remove(existing)
    tcMar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def _set_cell_borders(tc_elem, color_hex: str = "CBD5E1", sz: str = "4") -> None:
    """Set subtle cell borders via <w:tcBorders>."""
    tcPr = tc_elem.get_or_add_tcPr()
    for existing in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(existing)
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b_elem = OxmlElement(f"w:{side}")
        b_elem.set(qn("w:val"), "single")
        b_elem.set(qn("w:sz"), sz)
        b_elem.set(qn("w:space"), "0")
        b_elem.set(qn("w:color"), color_hex.lstrip("#"))
        tcBorders.append(b_elem)
    tcPr.append(tcBorders)


def _set_cell_valign(tc_elem, valign: str = "center") -> None:
    """Ensure vertical centering via <w:vAlign w:val="center"/>."""
    tcPr = tc_elem.get_or_add_tcPr()
    for existing in tcPr.findall(qn("w:vAlign")):
        tcPr.remove(existing)
    v_elem = OxmlElement("w:vAlign")
    v_elem.set(qn("w:val"), valign)
    tcPr.append(v_elem)


def _ensure_cell_has_paragraph(tc_elem) -> None:
    """
    The Last Paragraph Rule (ERR_DOCX_001):
    OpenXML standard requires that every cell (<w:tc>) must terminate with at least one paragraph (<w:p>).
    """
    paragraphs = tc_elem.findall(qn("w:p"))
    if not paragraphs:
        p = OxmlElement("w:p")
        tc_elem.append(p)


def _set_row_flags(tr_elem, is_header: bool = False, cant_split: bool = True) -> None:
    """Set <w:tblHeader> and <w:cantSplit> in <w:trPr>."""
    trPr = tr_elem.get_or_add_trPr()
    if cant_split and trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))
    if is_header and trPr.find(qn("w:tblHeader")) is None:
        trPr.append(OxmlElement("w:tblHeader"))


# ---------------------------------------------------------------------------
# Element Writers (Paragraph, Image, Table, Section)
# ---------------------------------------------------------------------------

def _write_paragraph(doc: DocumentType, para_data: dict[str, Any], anchor_para=None) -> Any:
    """Writes a styled paragraph with run-level formatting."""
    style_name = para_data.get("style_name", "Normal")
    style = _resolve_style(doc, style_name)
    if anchor_para is not None:
        para = anchor_para.insert_paragraph_before(style=style)
    else:
        para = doc.add_paragraph(style=style)

    _apply_paragraph_format(para, para_data)

    for run_data in para_data.get("runs", []):
        run = para.add_run(run_data.get("text", ""))
        _apply_run_format(run, run_data)

    if not para_data.get("runs") and para_data.get("text"):
        para.add_run(para_data["text"])

    return para


def _resolve_image_stream(img_data: dict[str, Any]) -> Tuple[Optional[io.BytesIO | str], Optional[str]]:
    """
    Resolves image binary from Base64 or local file path.
    Returns (stream_or_path, error_message).
    """
    # 1. Base64 data source
    raw_b64 = img_data.get("base64") or img_data.get("data")
    if raw_b64:
        try:
            if "," in raw_b64:
                raw_b64 = raw_b64.split(",", 1)[1]
            binary_data = base64.b64decode(raw_b64)
            return io.BytesIO(binary_data), None
        except Exception as e:
            return None, f"Invalid Base64 image payload: {e}"

    # 2. Local path source
    img_path = img_data.get("path") or img_data.get("src")
    if img_path:
        # Check absolute or relative paths
        candidate_paths = [
            img_path,
            os.path.abspath(img_path),
            os.path.join(os.getcwd(), img_path),
            os.path.join(os.path.dirname(os.path.abspath(__file__)), img_path),
            os.path.join(os.path.dirname(os.path.abspath(__file__)), "diagram_assets", os.path.basename(img_path)),
            os.path.join(os.path.dirname(os.path.abspath(__file__)), "specs", os.path.basename(img_path)),
        ]
        for p in candidate_paths:
            if os.path.isfile(p):
                return p, None

        return None, f"Image file not found: {img_path}"

    return None, "No path or Base64 data provided in image spec"


def _write_fallback_card(doc: DocumentType, title: str, error_msg: str, width_cm: float = 14.0) -> None:
    """Renders an elegant gray placeholder card when an image cannot be loaded."""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    tc = cell._tc
    _set_cell_shading(tc, "#F8FAFC")
    _set_cell_margins(tc, top=200, bottom=200, left=240, right=240)
    _set_cell_borders(tc, color_hex="94A3B8", sz="6")
    _set_cell_valign(tc, "center")

    for p in cell.paragraphs:
        tc.remove(p._p)

    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(f"📷 [Image Placeholder: {title}]\n")
    r1.bold = True
    r1.font.name = "Segoe UI"
    r1.font.size = Pt(10)
    r1.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    r2 = p.add_run(f"Status: {error_msg}")
    r2.italic = True
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    _ensure_cell_has_paragraph(tc)


def _embed_image_element(doc: DocumentType, img_data: dict[str, Any], anchor_para=None) -> None:
    """
    Embeds a real image into the DOCX with printable margin constraint enforcement (ERR_DOCX_005)
    and graceful fallback placeholder card.
    """
    stream_or_path, err = _resolve_image_stream(img_data)
    alt = img_data.get("alt_text") or img_data.get("caption") or img_data.get("image_name") or "Image"

    if err or not stream_or_path:
        _write_fallback_card(doc, alt, err or "Unknown image loading error")
        return

    # Calculate dimensions and enforce ERR_DOCX_005 (printable margin limit)
    current_sec = doc.sections[-1]
    pw = current_sec.page_width or Cm(21.0)
    lm = current_sec.left_margin or Cm(2.54)
    rm = current_sec.right_margin or Cm(2.54)
    usable_width_cm = pw.cm - lm.cm - rm.cm
    if usable_width_cm <= 0:
        usable_width_cm = 15.0

    raw_w = img_data.get("width_cm")
    raw_h = img_data.get("height_cm")
    target_w_cm: Optional[float] = float(raw_w) if raw_w is not None else None
    target_h_cm: Optional[float] = float(raw_h) if raw_h is not None else None

    try:
        # Inspect image aspect ratio with PIL
        with PILImage.open(stream_or_path) as pil_img:
            img_w_px, img_h_px = pil_img.size
            aspect = img_w_px / max(1, img_h_px)

        if stream_or_path and isinstance(stream_or_path, io.BytesIO):
            stream_or_path.seek(0)

        # Scale width to fit within printable margins
        if target_w_cm is None and target_h_cm is None:
            final_w_cm = min(usable_width_cm, 14.5)
            final_h_cm = final_w_cm / aspect
        elif target_w_cm is not None and target_h_cm is None:
            final_w_cm = min(target_w_cm, usable_width_cm)
            final_h_cm = final_w_cm / aspect
        elif target_w_cm is None and target_h_cm is not None:
            final_h_cm = target_h_cm
            final_w_cm = min(final_h_cm * aspect, usable_width_cm)
        elif target_w_cm is not None and target_h_cm is not None:
            final_w_cm = min(target_w_cm, usable_width_cm)
            final_h_cm = target_h_cm
        else:
            final_w_cm = min(usable_width_cm, 14.5)
            final_h_cm = final_w_cm / aspect

        # Add image paragraph
        align_str = str(img_data.get("alignment", "CENTER")).upper()
        alignment = _ALIGN_MAP.get(align_str, WD_ALIGN_PARAGRAPH.CENTER)

        if anchor_para is not None:
            p = anchor_para.insert_paragraph_before()
        else:
            p = doc.add_paragraph()

        p.alignment = alignment
        p.paragraph_format.space_before = Pt(img_data.get("space_before_pt", 6.0))
        p.paragraph_format.space_after = Pt(img_data.get("space_after_pt", 4.0))

        has_caption = bool(img_data.get("caption"))
        if has_caption:
            _set_keep_next(p)

        run = p.add_run()
        run.add_picture(stream_or_path, width=Cm(final_w_cm), height=Cm(final_h_cm))

        # Add caption if specified
        if has_caption:
            if anchor_para is not None:
                cp = anchor_para.insert_paragraph_before()
            else:
                cp = doc.add_paragraph()

            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_before = Pt(4.0)
            cp.paragraph_format.space_after = Pt(12.0)
            c_run = cp.add_run(img_data["caption"])
            c_run.font.name = "Times New Roman"
            c_run.font.size = Pt(10.0)
            c_run.italic = True
            c_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    except Exception as ex:
        _write_fallback_card(doc, alt, f"Failed to render image binary: {ex}")


def _write_table(doc: DocumentType, table_data: dict[str, Any], anchor_para=None) -> Any:
    """
    Constructs a production-grade OpenXML table supporting:
    - Merged cells with colspan & rowspan via automated grid mapping.
    - Table OpenXML Invariants: cantSplit, tblHeader, vAlign="center".
    - Last Paragraph Rule (ERR_DOCX_001).
    - Customizable header fills, zebra striping, and cell padding.
    """
    cells_flat = table_data.get("cells", [])
    
    # Calculate dimensions
    num_rows = table_data.get("rows")
    num_cols = table_data.get("cols")

    if not num_rows:
        num_rows = len(cells_flat)
    if not num_cols:
        # Determine maximum column span
        max_col_found = 0
        for r_cells in cells_flat:
            row_span_sum = sum(c.get("colspan", 1) for c in r_cells)
            max_col_found = max(max_col_found, row_span_sum, len(r_cells))
        num_cols = max(1, max_col_found)

    num_rows = max(1, num_rows)
    num_cols = max(1, num_cols)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = "Table Grid"

    # Virtual occupancy matrix to track merged cells
    occupied: list[list[bool]] = [[False for _ in range(num_cols)] for _ in range(num_rows)]

    # Iterate and build cells
    for r_idx, row_cells in enumerate(cells_flat):
        if r_idx >= num_rows:
            break

        row = table.rows[r_idx]
        is_header_row = (r_idx == 0) or any(c.get("is_header") for c in row_cells)
        _set_row_flags(row._tr, is_header=is_header_row, cant_split=True)

        col_cursor = 0
        for cell_data in row_cells:
            # Advance cursor to next unoccupied slot in this row
            while col_cursor < num_cols and occupied[r_idx][col_cursor]:
                col_cursor += 1
            if col_cursor >= num_cols:
                break

            c_idx = col_cursor
            colspan = max(1, int(cell_data.get("colspan", 1)))
            rowspan = max(1, int(cell_data.get("rowspan", 1)))

            # Clamp span to table boundaries
            end_r = min(num_rows - 1, r_idx + rowspan - 1)
            end_c = min(num_cols - 1, c_idx + colspan - 1)

            # Mark occupancy
            for rr in range(r_idx, end_r + 1):
                for cc in range(c_idx, end_c + 1):
                    occupied[rr][cc] = True

            # Merge if multi-cell span
            start_cell = table.cell(r_idx, c_idx)
            if end_r > r_idx or end_c > c_idx:
                end_cell = table.cell(end_r, end_c)
                target_cell = start_cell.merge(end_cell)
            else:
                target_cell = start_cell

            tc = target_cell._tc

            # Apply cell styling invariants
            pad_top = int(cell_data.get("padding_top_pt", 6.0) * 20)
            pad_bottom = int(cell_data.get("padding_bottom_pt", 6.0) * 20)
            pad_left = int(cell_data.get("padding_left_pt", 7.5) * 20)
            pad_right = int(cell_data.get("padding_right_pt", 7.5) * 20)
            _set_cell_margins(tc, top=pad_top, bottom=pad_bottom, left=pad_left, right=pad_right)
            _set_cell_borders(tc, color_hex=cell_data.get("border_color", "CBD5E1"), sz="4")
            _set_cell_valign(tc, valign=cell_data.get("valign", "center"))

            # Fill / shading
            fill_color = cell_data.get("fill_color")
            if not fill_color and is_header_row:
                fill_color = "#1F4E78"  # Enterprise Navy Header Fill
            if fill_color:
                _set_cell_shading(tc, fill_color)

            # Clear default paragraphs
            for p_elem in tc.findall(qn("w:p")):
                tc.remove(p_elem)

            # Populate cell content
            paragraphs_data = cell_data.get("paragraphs", [])
            if paragraphs_data:
                for para_data in paragraphs_data:
                    style_name = para_data.get("style_name", "Normal")
                    style = _resolve_style(doc, style_name)
                    p = target_cell.add_paragraph(style=style)
                    _apply_paragraph_format(p, para_data)

                    for run_data in para_data.get("runs", []):
                        run = p.add_run(run_data.get("text", ""))
                        _apply_run_format(run, run_data)
                        if is_header_row and fill_color in ("#1F4E78", "#2C3E50", "#000080", "#333399"):
                            if not run_data.get("font_color"):
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                run.bold = True

                    if not para_data.get("runs") and para_data.get("text"):
                        run = p.add_run(para_data["text"])
                        if is_header_row and fill_color in ("#1F4E78", "#2C3E50", "#000080", "#333399"):
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.bold = True
            elif cell_data.get("text") is not None:
                p = target_cell.add_paragraph()
                p.alignment = _ALIGN_MAP.get(str(cell_data.get("alignment", "LEFT")).upper(), WD_ALIGN_PARAGRAPH.LEFT)
                run = p.add_run(str(cell_data["text"]))
                if is_header_row and fill_color in ("#1F4E78", "#2C3E50", "#000080", "#333399"):
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.bold = True
                elif cell_data.get("font_bold"):
                    run.bold = True

            # Ensure ERR_DOCX_001
            _ensure_cell_has_paragraph(tc)

            col_cursor = end_c + 1

    # Apply Column Widths (Explicit or Heuristic Auto-Sizing)
    explicit_widths = table_data.get("col_widths_cm") or table_data.get("col_widths")
    if explicit_widths and isinstance(explicit_widths, list):
        col_widths_twips = [int(float(w) * 567.0) for w in explicit_widths]
    else:
        text_matrix = []
        for r_cells in cells_flat:
            row_txt = []
            for c in r_cells:
                txt = c.get("text", "")
                if not txt and c.get("paragraphs"):
                    txt = " ".join(p.get("text", "") for p in c["paragraphs"])
                row_txt.append(str(txt or ""))
            text_matrix.append(row_txt)

        try:
            sec = doc.sections[-1]
            pw = sec.page_width or Inches(8.5)
            lm = sec.left_margin or Inches(1.0)
            rm = sec.right_margin or Inches(1.0)
            avail_w = int((int(pw) - int(lm) - int(rm)) / 635)
        except Exception:
            avail_w = 9360

        col_widths_twips = calculate_heuristic_column_widths(text_matrix, total_width_twips=avail_w)

    if col_widths_twips:
        table.autofit = False
        tblGrid = table._tbl.tblGrid
        if tblGrid is not None:
            for cg in list(tblGrid):
                tblGrid.remove(cg)
            for w_tw in col_widths_twips:
                tblGrid.append(parse_xml(f'<w:gridCol {nsdecls("w")} w:w="{w_tw}"/>'))

        for row in table.rows:
            for c_idx, cell in enumerate(row.cells):
                if c_idx < len(col_widths_twips):
                    cell.width = Inches(col_widths_twips[c_idx] / 1440.0)

    # If anchor_para is provided, move table XML before anchor and remove anchor
    if anchor_para is not None:
        anchor_para._p.addprevious(table._tbl)
        parent = anchor_para._p.getparent()
        if parent is not None:
            parent.remove(anchor_para._p)

    return table


def _apply_section_break(doc: DocumentType, sec_data: dict[str, Any]) -> Any:
    """
    Inserts a new section break and configures orientation, margins, and header/footer.
    """
    new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    orient_str = str(sec_data.get("orientation", "portrait")).strip().lower()

    if orient_str == "landscape":
        new_sec.orientation = WD_ORIENT.LANDSCAPE
        # In python-docx, orientation switch requires explicitly swapping width and height
        w, h = new_sec.page_width, new_sec.page_height
        if w < h:
            new_sec.page_width = h
            new_sec.page_height = w
    else:
        new_sec.orientation = WD_ORIENT.PORTRAIT
        w, h = new_sec.page_width, new_sec.page_height
        if w > h:
            new_sec.page_width = h
            new_sec.page_height = w

    # Margins
    if sec_data.get("margin_top_cm") is not None:
        new_sec.top_margin = Cm(sec_data["margin_top_cm"])
    if sec_data.get("margin_bottom_cm") is not None:
        new_sec.bottom_margin = Cm(sec_data["margin_bottom_cm"])
    if sec_data.get("margin_left_cm") is not None:
        new_sec.left_margin = Cm(sec_data["margin_left_cm"])
    if sec_data.get("margin_right_cm") is not None:
        new_sec.right_margin = Cm(sec_data["margin_right_cm"])

    # Header configuration
    header_text = sec_data.get("header_text")
    if header_text:
        header_p = new_sec.header.paragraphs[0] if new_sec.header.paragraphs else new_sec.header.add_paragraph()
        header_p.text = str(header_text)
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if header_p.runs:
            header_p.runs[0].font.name = "Segoe UI"
            header_p.runs[0].font.size = Pt(8.5)
            header_p.runs[0].font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # Dynamic page number in footer
    if sec_data.get("footer_page_number", False):
        footer = new_sec.footer
        footer_p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        footer_p.text = ""
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        fldSimple = OxmlElement("w:fldSimple")
        fldSimple.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:ascii"), "Times New Roman")
        rFonts.set(qn("w:hAnsi"), "Times New Roman")
        rPr.append(rFonts)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "20")  # 10pt
        rPr.append(sz)
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(rPr)
        r.append(t)
        fldSimple.append(r)
        footer_p._p.append(fldSimple)

    return new_sec


# ---------------------------------------------------------------------------
# In-Place Template Patching Engine (Run-Merging Token Replacement)
# ---------------------------------------------------------------------------

def _replace_text_in_paragraph(paragraph, replacements: dict[str, str]) -> int:
    """
    Replaces {{KEY}} placeholders in a paragraph.
    Handles split runs by merging text into the first run of the match to preserve formatting.
    """
    total_replaced = 0
    full_text = paragraph.text
    if not full_text:
        return 0

    for key, val in replacements.items():
        if key not in full_text:
            continue

        runs = paragraph.runs
        if not runs:
            continue

        # Simple single-run optimization
        replaced_in_single_run = False
        for r in runs:
            if key in r.text:
                r.text = r.text.replace(key, val)
                replaced_in_single_run = True
                total_replaced += 1

        if replaced_in_single_run:
            full_text = paragraph.text
            continue

        # Multi-run split resolution
        # Build character run spans: [(start_idx, end_idx, run_obj)]
        spans = []
        cur_idx = 0
        for r in runs:
            start = cur_idx
            end = cur_idx + len(r.text)
            spans.append((start, end, r))
            cur_idx = end

        # Locate key in full string
        k_idx = full_text.find(key)
        while k_idx != -1:
            k_end = k_idx + len(key)
            # Find runs overlapping with [k_idx, k_end]
            matching_runs = []
            for s_start, s_end, r_obj in spans:
                if s_start < k_end and s_end > k_idx:
                    matching_runs.append((s_start, s_end, r_obj))

            if matching_runs:
                # First run gets the replacement text
                first_start, first_end, first_run = matching_runs[0]
                prefix = first_run.text[:max(0, k_idx - first_start)]
                first_run.text = prefix + val

                # Subsequent overlapping runs have the matched slice stripped
                for other_start, other_end, other_run in matching_runs[1:]:
                    if other_end <= k_end:
                        other_run.text = ""
                    else:
                        suffix_offset = max(0, k_end - other_start)
                        other_run.text = other_run.text[suffix_offset:]

                total_replaced += 1

            full_text = paragraph.text
            k_idx = full_text.find(key)

    return total_replaced


def patch_docx_template(
    template_path: str,
    output_path: str,
    text_replacements: dict[str, str] | None = None,
    block_replacements: dict[str, Any] | None = None,
) -> str:
    """
    In-Place Template Patching:
    Updates an existing Word template document without clearing its body,
    preserving 100% of the Cover Page, shapes, styles, and DrawingML objects.

    Args:
        template_path:      Path to the source .docx template.
        output_path:        Destination path for the patched .docx.
        text_replacements:  Dict of {"{{KEY}}": "Value"} to replace across body, tables, and headers/footers.
        block_replacements: Dict of {"{{BLOCK:ID}}": element_spec} where element_spec can be:
                            - {"type": "table", "rows": ..., "cols": ..., "cells": ...}
                            - {"type": "image", "path": ..., "caption": ...}
                            - {"type": "paragraph", "text": ..., "runs": ...}
                            - list of element specs

    Returns:
        Absolute path to the saved output file.
    """
    if not os.path.isfile(template_path):
        raise FileNotFoundError(f"Template DOCX not found: {template_path}")

    doc = Document(template_path)
    text_rep = text_replacements or {}
    block_rep = block_replacements or {}

    # 1. Text Replacements in Paragraphs
    if text_rep:
        for p in doc.paragraphs:
            _replace_text_in_paragraph(p, text_rep)

        # In Tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        _replace_text_in_paragraph(p, text_rep)

        # In Headers & Footers
        for sec in doc.sections:
            for p in sec.header.paragraphs:
                _replace_text_in_paragraph(p, text_rep)
            for table in sec.header.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            _replace_text_in_paragraph(p, text_rep)

            for p in sec.footer.paragraphs:
                _replace_text_in_paragraph(p, text_rep)
            for table in sec.footer.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            _replace_text_in_paragraph(p, text_rep)

    # 2. Block Replacements (Anchors in Document Body)
    if block_rep:
        for block_key, block_spec in block_rep.items():
            # Search for anchor paragraph
            for p in list(doc.paragraphs):
                if block_key in p.text:
                    if isinstance(block_spec, list):
                        for el in block_spec:
                            el_type = el.get("type", "paragraph")
                            if el_type == "paragraph":
                                _write_paragraph(doc, el, anchor_para=p)
                            elif el_type == "image":
                                _embed_image_element(doc, el, anchor_para=p)
                            elif el_type == "table":
                                _write_table(doc, el, anchor_para=p)
                        # Remove anchor paragraph
                        parent = p._p.getparent()
                        if parent is not None:
                            parent.remove(p._p)
                    elif isinstance(block_spec, dict):
                        el_type = block_spec.get("type", "paragraph")
                        if el_type == "table":
                            _write_table(doc, block_spec, anchor_para=p)
                        elif el_type == "image":
                            _embed_image_element(doc, block_spec, anchor_para=p)
                            parent = p._p.getparent()
                            if parent is not None:
                                parent.remove(p._p)
                        elif el_type == "paragraph":
                            _write_paragraph(doc, block_spec, anchor_para=p)
                            parent = p._p.getparent()
                            if parent is not None:
                                parent.remove(p._p)
                    break

    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    doc.save(output_path)
    return os.path.abspath(output_path)


# ---------------------------------------------------------------------------
# Public API: Generative Mode (Snapshot → DOCX)
# ---------------------------------------------------------------------------

def write_docx(
    snapshot: dict[str, Any],
    output_path: str,
    template_path: str | None = None,
) -> str:
    """
    Write a DOCX file from a JSON snapshot.

    Args:
        snapshot:      Dict produced by docx_reader.read_docx() or declarative generator.
        output_path:   Destination .docx path.
        template_path: Optional path to original DOCX used as style anchor.

    Returns:
        Absolute path of the written file.
    """
    if template_path and os.path.isfile(template_path):
        doc = Document(template_path)
        # Clear body elements to prepare for generative reconstruction, keep styles & sections
        body = doc.element.body
        for child in list(body):
            tag = child.tag
            if tag not in (qn("w:sectPr"),):
                body.remove(child)
    else:
        doc = Document()

    # Initial section layout
    sections_data = snapshot.get("sections", [])
    if sections_data:
        sec = doc.sections[0]
        sd = sections_data[0]
        if sd.get("orientation", "").lower() == "landscape":
            sec.orientation = WD_ORIENT.LANDSCAPE
            w = sec.page_width or Inches(8.5)
            h = sec.page_height or Inches(11.0)
            if int(w) < int(h):
                sec.page_width = h
                sec.page_height = w
        if sd.get("page_width_cm"):
            sec.page_width = Cm(sd["page_width_cm"])
        if sd.get("page_height_cm"):
            sec.page_height = Cm(sd["page_height_cm"])
        if sd.get("margin_top_cm") is not None:
            sec.top_margin = Cm(sd["margin_top_cm"])
        if sd.get("margin_bottom_cm") is not None:
            sec.bottom_margin = Cm(sd["margin_bottom_cm"])
        if sd.get("margin_left_cm") is not None:
            sec.left_margin = Cm(sd["margin_left_cm"])
        if sd.get("margin_right_cm") is not None:
            sec.right_margin = Cm(sd["margin_right_cm"])

    # Write body elements in sequential stream
    for element in snapshot.get("body", []):
        elem_type = element.get("type")
        if elem_type == "paragraph":
            _write_paragraph(doc, element)
        elif elem_type == "table":
            _write_table(doc, element)
        elif elem_type == "image":
            _embed_image_element(doc, element)
        elif elem_type == "section_break":
            _apply_section_break(doc, element)

    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    doc.save(output_path)
    return os.path.abspath(output_path)


def write_docx_from_json_file(
    json_path: str,
    output_path: str,
    template_path: str | None = None,
) -> str:
    """Load a JSON snapshot from file and write it to a DOCX."""
    if not os.path.isfile(json_path):
        raise FileNotFoundError(f"JSON snapshot not found: {json_path}")

    with open(json_path, "r", encoding="utf-8") as f:
        snapshot = json.load(f)

    return write_docx(snapshot, output_path, template_path)


# ---------------------------------------------------------------------------
# Surgical Heading & Diagram Injection API (Backward Compatible)
# ---------------------------------------------------------------------------

def inject_content_into_docx(
    source_docx_path: str,
    spec: dict[str, Any],
    output_path: str | None = None,
) -> str:
    """
    Surgically inject paragraphs, runs, formatting, tables, and images into an existing DOCX
    at a targeted location (e.g. after a heading or placeholder).
    """
    if not os.path.isfile(source_docx_path):
        raise FileNotFoundError(f"Source DOCX not found: {source_docx_path}")

    doc = Document(source_docx_path)

    target_idx: int | None = None
    is_placeholder_mode = False
    placeholder = spec.get("placeholder")
    target_heading = spec.get("target_heading")

    if placeholder:
        for i, p in enumerate(doc.paragraphs):
            if placeholder in p.text:
                target_idx = i
                is_placeholder_mode = True
                break

    if target_idx is None and target_heading:
        for i, p in enumerate(doc.paragraphs):
            if target_heading in p.text:
                target_idx = i
                break
    elif target_idx is None and spec.get("target_element_index") is not None:
        target_idx = int(spec["target_element_index"])

    if target_idx is None:
        target_desc = f"placeholder '{placeholder}'" if placeholder else f"heading '{target_heading}'"
        raise ValueError(f"Target {target_desc} not found in document")

    paragraphs_to_remove = []

    if is_placeholder_mode:
        anchor_para = doc.paragraphs[target_idx]
        paragraphs_to_remove.append(anchor_para)
    else:
        replace_empty = spec.get("replace_empty_after", False)
        curr_idx = target_idx + 1

        if replace_empty:
            while curr_idx < len(doc.paragraphs):
                p_next = doc.paragraphs[curr_idx]
                drawings = p_next._element.findall(".//" + qn("w:drawing"))
                if not p_next.text.strip() and not drawings:
                    paragraphs_to_remove.append(p_next)
                    curr_idx += 1
                else:
                    break

        anchor_para = doc.paragraphs[curr_idx] if curr_idx < len(doc.paragraphs) else None

    # Inject elements
    for el in spec.get("elements", []):
        elem_type = el.get("type", "paragraph")
        if elem_type == "paragraph":
            style_name = el.get("style_name", "Normal")
            style = _resolve_style(doc, style_name)
            if anchor_para is not None:
                p = anchor_para.insert_paragraph_before(style=style)
            else:
                p = doc.add_paragraph(style=style)

            _apply_paragraph_format(p, el)

            # Add images if specified
            for img_info in el.get("images", []):
                _embed_image_element(doc, img_info, anchor_para=p)
                if img_info.get("gap", True):
                    p.add_run("   ")

            # Add runs
            for run_data in el.get("runs", []):
                r = p.add_run(run_data.get("text", ""))
                _apply_run_format(r, run_data)

            if not el.get("runs") and not el.get("images") and el.get("text"):
                p.add_run(el["text"])

        elif elem_type == "table":
            _write_table(doc, el, anchor_para=anchor_para)
        elif elem_type == "image":
            _embed_image_element(doc, el, anchor_para=anchor_para)

    # Remove placeholder or empty paragraphs
    for p_rem in paragraphs_to_remove:
        parent = p_rem._element.getparent()
        if parent is not None:
            parent.remove(p_rem._element)

    dest = os.path.abspath(output_path or source_docx_path)
    if dest == os.path.abspath(source_docx_path):
        bak_path = source_docx_path + ".bak"
        shutil.copy2(source_docx_path, bak_path)

    out_dir = os.path.dirname(dest)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    doc.save(dest)
    return dest


def inject_diagram_into_docx(
    source_docx_path: str | None = None,
    diagram_spec: dict[str, Any] | None = None,
    output_path: str | None = None,
    *,
    docx_path: str | None = None,
    png_path: str | None = None,
    heading: str | None = None,
    placeholder: str | None = None,
    caption: str | dict[str, Any] | None = None,
    width_cm: float | None = None,
    height_cm: float | None = None,
    max_height_cm: float | None = None,
) -> str:
    """
    Convenience wrapper to inject a diagram into a DOCX at a target location with keepNext chaining.
    """
    doc_target = source_docx_path or docx_path
    if not doc_target:
        raise ValueError("Missing required DOCX path (source_docx_path or docx_path)")

    if diagram_spec is None:
        diagram_spec = {}
    else:
        diagram_spec = dict(diagram_spec)

    if png_path:
        diagram_spec.setdefault("image_path", png_path)
    if heading:
        diagram_spec.setdefault("target_heading", heading)
    if placeholder:
        diagram_spec.setdefault("placeholder", placeholder)
    if caption:
        diagram_spec.setdefault("caption_template", caption)
    if width_cm is not None:
        diagram_spec.setdefault("width_cm", width_cm)
    if height_cm is not None:
        diagram_spec.setdefault("height_cm", height_cm)
    elif max_height_cm is not None:
        diagram_spec.setdefault("height_cm", max_height_cm)

    image_path = diagram_spec.get("image_path")
    if not image_path or not os.path.isfile(image_path):
        raise FileNotFoundError(f"Diagram image not found: {image_path}")

    width_cm_val = diagram_spec.get("width_cm", 14.0)
    height_cm_val = diagram_spec.get("height_cm")

    cap_template = diagram_spec.get("caption_template")
    caption_text = ""
    description_text = ""
    cap_style: dict[str, Any] = {}
    desc_style: dict[str, Any] = {}

    if isinstance(cap_template, str):
        caption_text = cap_template
    elif isinstance(cap_template, dict):
        caption_text = cap_template.get("caption", "")
        description_text = cap_template.get("description", "")
        cap_style = cap_template.get("style", {})
        desc_style = cap_template.get("description_style", {})

    has_caption = bool(caption_text.strip())
    has_description = bool(description_text)

    elements: list[dict[str, Any]] = []

    # 1. Image paragraph
    elements.append({
        "type": "paragraph",
        "alignment": "CENTER",
        "space_before_pt": 6,
        "space_after_pt": 4,
        "keep_next": has_caption or has_description,
        "images": [{
            "path": image_path,
            "width_cm": width_cm_val,
            "height_cm": height_cm_val,
        }],
    })

    # 2. Caption paragraph
    if has_caption:
        elements.append({
            "type": "paragraph",
            "alignment": cap_style.get("caption_alignment", "CENTER"),
            "space_before_pt": cap_style.get("space_before_pt", 4),
            "space_after_pt": cap_style.get("space_after_pt", 6 if has_description else 12),
            "keep_next": has_description,
            "runs": [{
                "text": caption_text,
                "bold": cap_style.get("caption_bold", False),
                "italic": cap_style.get("caption_italic", True),
                "font_name": cap_style.get("font_name", "Calibri"),
                "font_size": cap_style.get("caption_font_size", 10.0),
                "font_color": cap_style.get("caption_color", "#595959"),
            }],
        })

    # 3. Description paragraph(s)
    if has_description:
        desc_paragraphs = [description_text] if isinstance(description_text, str) else list(description_text)
        for d_idx, d_text in enumerate(desc_paragraphs):
            is_last = (d_idx == len(desc_paragraphs) - 1)
            elements.append({
                "type": "paragraph",
                "alignment": desc_style.get("alignment", "JUSTIFY"),
                "space_before_pt": desc_style.get("space_before_pt", 3),
                "space_after_pt": desc_style.get("space_after_pt", 12 if is_last else 4),
                "left_indent_pt": desc_style.get("left_indent_pt", 0),
                "keep_next": not is_last,
                "runs": [{
                    "text": d_text,
                    "bold": False,
                    "italic": False,
                    "font_name": desc_style.get("font_name", "Calibri"),
                    "font_size": desc_style.get("font_size", 10.5),
                    "font_color": desc_style.get("font_color", "#262626"),
                }],
            })

    injection_spec = {
        "placeholder": diagram_spec.get("placeholder"),
        "target_heading": diagram_spec.get("target_heading"),
        "target_element_index": diagram_spec.get("target_element_index"),
        "replace_empty_after": diagram_spec.get("replace_empty_after", False),
        "elements": elements,
    }

    return inject_content_into_docx(doc_target, injection_spec, output_path)
