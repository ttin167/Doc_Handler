"""
pptx_reader.py — OpenXML PowerPoint Reader & AST Snapshot Parser (v3.0).

Reads .pptx presentations and reconstructs a structured JSON Abstract Syntax Tree (AST):
- Presentation dimensions & aspect ratio
- Individual slide metadata, layout index, and notes
- Shapes (Text frames, native tables, picture dimensions, and coordinates)
"""

from __future__ import annotations

import json
import os
from typing import Any, Dict, List, Optional

import pptx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def read_pptx(pptx_path: str) -> Dict[str, Any]:
    """
    Parses a PowerPoint presentation into a rich JSON-serializable AST dictionary.
    """
    if not os.path.isfile(pptx_path):
        raise FileNotFoundError(f"PowerPoint file not found: {pptx_path}")

    prs = Presentation(pptx_path)
    width_in = prs.slide_width.inches
    height_in = prs.slide_height.inches

    aspect_ratio = "16:9" if abs(width_in / height_in - 16 / 9) < 0.1 else "4:3"

    slides_ast: List[Dict[str, Any]] = []

    for idx, slide in enumerate(prs.slides, start=1):
        slide_title = ""
        shapes_data: List[Dict[str, Any]] = []

        # Check for standard slide title shape
        if slide.shapes.title and slide.shapes.title.text_frame:
            slide_title = slide.shapes.title.text_frame.text.strip()

        for s in slide.shapes:
            shape_type_str = str(s.shape_type)
            shape_info: Dict[str, Any] = {
                "name": s.name,
                "type": shape_type_str,
                "left_in": round(s.left.inches, 3) if s.left else 0.0,
                "top_in": round(s.top.inches, 3) if s.top else 0.0,
                "width_in": round(s.width.inches, 3) if s.width else 0.0,
                "height_in": round(s.height.inches, 3) if s.height else 0.0,
            }

            # 1. Text Box / Text Frame
            if s.has_text_frame:
                paragraphs = []
                for p in s.text_frame.paragraphs:
                    p_text = p.text.strip()
                    if p_text:
                        paragraphs.append({
                            "text": p_text,
                            "level": p.level,
                        })
                shape_info["paragraphs"] = paragraphs
                if not slide_title and paragraphs:
                    slide_title = paragraphs[0]["text"]

            # 2. Native Table
            elif s.has_table:
                tbl = s.table
                table_rows = []
                for row in tbl.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    table_rows.append(row_cells)
                shape_info["table"] = {
                    "num_rows": len(tbl.rows),
                    "num_cols": len(tbl.columns),
                    "rows": table_rows,
                }

            # 3. Picture / Diagram
            elif s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                shape_info["is_picture"] = True
                shape_info["content_type"] = getattr(s.image, "content_type", "image")

            shapes_data.append(shape_info)

        # Presenter Notes
        notes_text = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()

        slides_ast.append({
            "slide_number": idx,
            "title": slide_title,
            "shapes_count": len(shapes_data),
            "shapes": shapes_data,
            "notes": notes_text,
        })

    return {
        "file_path": os.path.abspath(pptx_path),
        "dimensions": {
            "width_inches": width_in,
            "height_inches": height_in,
            "aspect_ratio": aspect_ratio,
        },
        "total_slides": len(slides_ast),
        "slides": slides_ast,
    }


def read_pptx_to_json(pptx_path: str, output_json: Optional[str] = None) -> str:
    """
    Parses a PowerPoint file and saves the AST to a JSON file.
    """
    ast = read_pptx(pptx_path)
    if not output_json:
        base, _ = os.path.splitext(pptx_path)
        output_json = f"{base}_ast.json"

    with open(output_json, "w", encoding="utf-8") as f:
        json.dump(ast, f, indent=2, ensure_ascii=False)

    return os.path.abspath(output_json)
