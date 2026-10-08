"""
pptx_engine.py — High-Level Coordinator, Markdown-to-PPTX Compiler & Template Engine.

Key Capabilities:
1. Markdown Slide Compiler: Converts Markdown divided by '---' into 16:9 Widescreen PowerPoint decks.
2. Auto-Detection of Diagram Slides: Parses '![caption](path.png)' and locks aspect ratios.
3. Native Table Generation: Converts Markdown tables '| col | col |' into OpenXML native tables.
4. In-Place Diagram Injection: Inserts technical diagrams into existing presentation files.
"""

from __future__ import annotations

import os
import re
import sys
from typing import Any, Dict, List, Optional, Tuple

from pptx import Presentation
from pptx.util import Inches, Pt

try:
    from .pptx_writer import (
        write_pptx_from_spec,
        add_title_slide,
        add_chapter_slide,
        add_content_slide,
        add_diagram_slide,
        add_split_diagram_slide,
        add_dual_diagram_slide,
        add_metrics_summary_slide,
        add_table_slide,
        add_two_column_slide,
        add_kpi_cards_slide,
        init_presentation,
        safe_save_pptx,
        get_theme,
        sanitize_openxml_presentation,
        set_paragraph_text_clean,
    )
except (ImportError, ValueError):
    from pptx_writer import (  # type: ignore
        write_pptx_from_spec,
        add_title_slide,
        add_chapter_slide,
        add_content_slide,
        add_diagram_slide,
        add_split_diagram_slide,
        add_dual_diagram_slide,
        add_metrics_summary_slide,
        add_table_slide,
        add_two_column_slide,
        add_kpi_cards_slide,
        init_presentation,
        safe_save_pptx,
        get_theme,
        sanitize_openxml_presentation,
        set_paragraph_text_clean,
    )

try:
    from .pptx_math import add_math_text_to_paragraph, parse_inline_math
except (ImportError, ValueError):
    from pptx_math import add_math_text_to_paragraph, parse_inline_math  # type: ignore


def parse_markdown_to_slides_spec(md_text: str, default_theme: str = "corporate_blue") -> Dict[str, Any]:
    """
    Parses a Markdown slide deck string into a declarative Slide Spec dictionary.
    Delimiters: '---' or '==='.
    Auto-detects:
    - Bullets + 1 image  -> split_diagram slide (card left + diagram right)
    - Bullets + 2 images -> dual_diagram slide (card left + 2 stacked diagrams)
    - 1 image only       -> diagram slide (centered full-width)
    - Tables             -> table slide (OpenXML native styled table)
    """
    raw_slides = re.split(r'\n(?:\-{3,}|\={3,})\n', md_text.strip())
    slides_data: List[Dict[str, Any]] = []

    for idx, raw_slide in enumerate(raw_slides):
        lines = [line.rstrip() for line in raw_slide.strip().split('\n') if line.strip()]
        if not lines:
            continue

        slide_type = "content"
        title = ""
        subtitle = None
        banner = None
        bullets: List[Tuple[str, int]] = []
        images: List[Tuple[str, str]] = []
        table_rows: List[List[str]] = []
        notes = ""

        # Check for title
        first_line = lines[0]
        if first_line.startswith("# "):
            title = first_line[2:].strip()
            lines = lines[1:]
            if idx == 0:
                slide_type = "title"
        elif first_line.startswith("## "):
            title = first_line[3:].strip()
            lines = lines[1:]
        elif first_line.startswith("### "):
            title = first_line[4:].strip()
            lines = lines[1:]

        # Parse remaining lines
        for line in lines:
            # Notes syntax: <!-- note: ... -->
            if "<!-- note:" in line:
                m_note = re.search(r'<!--\s*note:\s*(.*?)\s*-->', line, re.IGNORECASE)
                if m_note:
                    notes = m_note.group(1).strip()
                continue

            # Banner syntax: <!-- banner: ... -->
            if "<!-- banner:" in line:
                m_b = re.search(r'<!--\s*banner:\s*(.*?)\s*-->', line, re.IGNORECASE)
                if m_b:
                    banner = m_b.group(1).strip()
                continue

            # Subtitle syntax: > Subtitle
            if line.startswith("> "):
                subtitle = line[2:].strip()
                continue

            # Image / Diagram syntax: ![Caption](path/to/img.png)
            m_img = re.search(r'!\[(.*?)\]\((.*?)\)', line)
            if m_img:
                cap, path = m_img.group(1), m_img.group(2)
                images.append((path, cap))
                continue

            # Markdown Table Syntax: | Col 1 | Col 2 |
            if line.startswith("|") and line.endswith("|"):
                # Skip divider line like |---|---|
                if re.match(r'^\|[\s\-:|]+\|$', line):
                    continue
                cells = [c.strip() for c in line.strip("|").split("|")]
                table_rows.append(cells)
                continue

            # Bullets
            m_bullet = re.match(r'^(\s*)([\*\-\+])\s+(.*)$', line)
            if m_bullet:
                indent = len(m_bullet.group(1))
                level = min(indent // 2, 2)
                text = m_bullet.group(3).strip()
                bullets.append((text, level))
                continue

            # Regular text line
            if not title:
                title = line
            elif not subtitle and idx == 0:
                subtitle = line
            else:
                bullets.append((line, 0))

        # Determine final slide spec
        if table_rows and len(table_rows) >= 2:
            slide_spec = {
                "type": "table",
                "title": title or "Table Overview",
                "subtitle": subtitle,
                "banner": banner,
                "headers": table_rows[0],
                "rows": table_rows[1:],
                "notes": notes,
            }
        elif len(images) >= 2:
            slide_spec = {
                "type": "dual_diagram",
                "title": title or "Comparative Analysis",
                "subtitle": subtitle,
                "banner": banner,
                "bullets": bullets,
                "img1_path": images[0][0],
                "cap1": images[0][1],
                "img2_path": images[1][0],
                "cap2": images[1][1],
                "notes": notes,
            }
        elif len(images) == 1:
            if bullets:
                slide_spec = {
                    "type": "split_diagram",
                    "title": title or "Technical Specification",
                    "subtitle": subtitle,
                    "banner": banner,
                    "bullets": bullets,
                    "image_path": images[0][0],
                    "caption": images[0][1] or None,
                    "notes": notes,
                }
            else:
                slide_spec = {
                    "type": "diagram",
                    "title": title or "Diagram Overview",
                    "subtitle": subtitle,
                    "banner": banner,
                    "image_path": images[0][0],
                    "caption": images[0][1] or None,
                    "notes": notes,
                }
        elif slide_type == "title":
            slide_spec = {
                "type": "title",
                "title": title or "Untitled Presentation",
                "subtitle": subtitle,
                "notes": notes,
            }
        else:
            slide_spec = {
                "type": "content",
                "title": title or "Slide Title",
                "subtitle": subtitle,
                "banner": banner,
                "bullets": bullets,
                "notes": notes,
            }

        slides_data.append(slide_spec)

    return {
        "theme": default_theme,
        "slides": slides_data,
    }


def compile_markdown_file_to_pptx(
    md_file_path: str,
    output_pptx_path: str,
    theme_name: str = "corporate_blue",
    template_path: Optional[str] = None,
) -> str:
    """
    Reads a Markdown file and compiles it into a .pptx presentation.
    """
    if not os.path.isfile(md_file_path):
        raise FileNotFoundError(f"Markdown slides file not found: {md_file_path}")

    with open(md_file_path, "r", encoding="utf-8") as f:
        md_content = f.read()

    spec = parse_markdown_to_slides_spec(md_content, default_theme=theme_name)
    return write_pptx_from_spec(
        spec=spec,
        output_path=output_pptx_path,
        theme_name=theme_name,
        template_path=template_path,
    )


def compile_markdown_to_presentation(
    md_text: str,
    output_pptx_path: str,
    theme_name: str = "corporate_blue",
    template_path: Optional[str] = None,
) -> str:
    """
    Directly compiles a Markdown slide deck string into a .pptx presentation.
    """
    spec = parse_markdown_to_slides_spec(md_text, default_theme=theme_name)
    return write_pptx_from_spec(
        spec=spec,
        output_path=output_pptx_path,
        theme_name=theme_name,
        template_path=template_path,
    )


def insert_diagram_into_pptx(
    pptx_path: str,
    image_path: str,
    slide_index: Optional[int] = None,
    title: Optional[str] = None,
    caption: Optional[str] = None,
    output_path: Optional[str] = None,
    theme_name: str = "corporate_blue",
) -> str:
    """
    Inserts a technical diagram into a slide within an existing PowerPoint presentation.
    If slide_index is not provided, appends a new diagram slide.
    """
    if not os.path.isfile(pptx_path):
        raise FileNotFoundError(f"Target PPTX not found: {pptx_path}")
    if not os.path.isfile(image_path):
        raise FileNotFoundError(f"Diagram image not found: {image_path}")

    prs = Presentation(pptx_path)

    if slide_index is not None and 1 <= slide_index <= len(prs.slides):
        # Insert picture into existing slide
        slide = prs.slides[slide_index - 1]
        theme = get_theme(theme_name)
        # Position in lower half or main body
        slide.shapes.add_picture(
            image_path,
            Inches(1.0),
            Inches(2.0),
            width=Inches(11.333),
        )
    else:
        # Append a new diagram slide
        add_diagram_slide(
            prs=prs,
            title=title or "System Architecture Diagram",
            image_path=image_path,
            caption=caption,
            theme_name=theme_name,
        )

    out = output_path or pptx_path
    return safe_save_pptx(prs, out)
