"""
test_excel_core.py — Automated Test Suite for UniversalExcelEngine & Excel Invariants.

Verifies:
- ERR_XLSX_001: Prototype Row Style Cloning (Font, Fill, Border, Alignment, NumberFormat).
- ERR_XLSX_002: Dynamic Formula Shifting (Regex AST expands ranges and shifts downstream cells).
- ERR_XLSX_003: Semantic Anchor Discovery (Keyword header scanning).
- ERR_XLSX_004: Safe Merged-Cell Introspection & Border Synchronization.
- ERR_XLSX_005: 2D Freeze Panes & Auto-Fit Row Height / Column Width.
- ERR_XLSX_006: DrawingML & Template Preservation.
- ERR_XLSX_007: Identifier '@' string formatting.
"""

import os
import sys
import unittest
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from xlsx_writer import (
    UniversalExcelEngine,
    mutate_template_excel,
    shift_formula_string,
    safe_set_cell,
)


class TestExcelCoreEngine(unittest.TestCase):
    def setUp(self):
        self.output_dir = os.path.join(BASE_DIR, "scratch", "test_excel_out")
        os.makedirs(self.output_dir, exist_ok=True)

    def test_regex_ast_formula_shifting(self):
        """Verify Regex AST shifts ranges across insertion point and downstream cells (ERR_XLSX_002)."""
        # Spanning range: insert 10 rows at row 15
        self.assertEqual(shift_formula_string("=SUM(F10:F25)", insert_row=15, delta=10), "=SUM(F10:F35)")
        # Range above insert point: unaffected
        self.assertEqual(shift_formula_string("=SUM(A1:A5)", insert_row=15, delta=10), "=SUM(A1:A5)")
        # Downstream single cell: shifted
        self.assertEqual(shift_formula_string("=F26*2", insert_row=15, delta=10), "=F36*2")
        # Formula with absolute markers
        self.assertEqual(
            shift_formula_string('=COUNTIF($G$10:$G$25, "Success") + F26', insert_row=15, delta=10),
            '=COUNTIF($G$10:$G$35, "Success") + F36'
        )

    def test_universal_excel_mutation_lifecycle(self):
        """Build mock template, mutate dynamically with row expansion, verify styles & formulas."""
        template_path = os.path.join(self.output_dir, "mock_template.xlsx")
        output_path = os.path.join(self.output_dir, "mock_mutated.xlsx")

        # 1. Create a mock template workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Test Cases"

        # Title
        ws["A1"] = "PROJECT: {{PROJECT_TITLE}}"
        ws["A1"].font = Font(name="Segoe UI", size=14, bold=True)

        # Header Sentinel row (Row 5)
        headers = ["Test ID", "Description", "Result", "Score"]
        for c_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=5, column=c_idx, value=h_text)
            cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFFFF")
            cell.fill = PatternFill(fill_type="solid", start_color="FF1F4E78", end_color="FF1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Prototype Data Row (Row 6)
        proto_data = ["TC_000", "Sample Prototype Description", "Success", 10]
        for c_idx, val in enumerate(proto_data, start=1):
            cell = ws.cell(row=6, column=c_idx, value=val)
            cell.font = Font(name="Segoe UI", size=10, bold=False)
            cell.fill = PatternFill(fill_type="solid", start_color="FFF8FAFC", end_color="FFF8FAFC")
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(
                top=Side(style="thin", color="FFCBD5E1"),
                bottom=Side(style="thin", color="FFCBD5E1"),
                left=Side(style="thin", color="FFCBD5E1"),
                right=Side(style="thin", color="FFCBD5E1"),
            )

        # Summary Row (Row 7)
        ws["C7"] = "Total Score"
        ws["C7"].font = Font(name="Segoe UI", size=11, bold=True)
        ws["D7"] = "=SUM(D6:D6)"
        ws["D7"].font = Font(name="Segoe UI", size=11, bold=True)

        # Merged range test
        ws.merge_cells("A10:C11")
        ws["A10"] = "Important Notice Block"

        wb.save(template_path)

        # 2. Mutate template using mutate_template_excel
        mutation_spec = {
            "sheets": {
                "Test Cases": {
                    "cell_updates": {
                        "A1": "PROJECT: ZSCORT Automated Test Matrix",
                    },
                    "table_expansions": [
                        {
                            "anchor_keyword": "Test ID",  # Automatically discovers Row 5
                            "prototype_row": 6,
                            "rows": [
                                ["001", "Verify OAuth2 Authentication Flow\nWith Refresh Token", "Success", 95],
                                ["002", "Verify High-Speed PDF Conversion", "Success", 100],
                                ["003", "Verify OpenXML Table Repair Invariants", "Success", 98],
                            ],
                            "id_columns": [0],  # "001" must preserve leading zeros via '@'
                            "auto_fit": True,
                        }
                    ],
                    "freeze_panes": "B6",
                }
            }
        }

        res_path = mutate_template_excel(template_path, output_path, mutation_spec)
        self.assertTrue(os.path.isfile(res_path))

        # 3. Inspect generated mutated workbook
        m_wb = openpyxl.load_workbook(res_path, data_only=False)
        m_ws = m_wb["Test Cases"]

        # Check title update
        self.assertEqual(m_ws["A1"].value, "PROJECT: ZSCORT Automated Test Matrix")

        # Check rows expansion: inserted 3 rows at row 6, so previous row 6 shifted to row 9, summary at row 10
        # Rows 6, 7, 8 are new data
        self.assertEqual(m_ws.cell(row=6, column=1).value, "001")
        # Identifier Number Format '@' (ERR_XLSX_007)
        self.assertEqual(m_ws.cell(row=6, column=1).number_format, "@")

        # Check prototype style cloning (ERR_XLSX_001)
        self.assertEqual(m_ws.cell(row=6, column=1).font.name, "Segoe UI")
        self.assertIsNotNone(m_ws.cell(row=6, column=1).border.top)

        # Check formula shifting in summary row (Row 10) (ERR_XLSX_002)
        # Original: =SUM(D6:D6), shifted by 3 rows: =SUM(D6:D9)
        summary_formula = m_ws.cell(row=10, column=4).value
        self.assertEqual(summary_formula, "=SUM(D6:D9)")

        # Check Freeze Panes (ERR_XLSX_005)
        self.assertEqual(m_ws.freeze_panes, "B6")

        # Check Merged Range shifted down (Original A10:C11 shifted by 3 -> A13:C14)
        merged_coords = [str(r.coord) for r in m_ws.merged_cells.ranges]
        self.assertIn("A13:C14", merged_coords)

    def test_multi_axis_and_cross_sheet_formula_shifting(self):
        """Verify Regex AST shifts across rows, columns, and preserves sheet prefixes (ERR_XLSX_002)."""
        # Spanning row range
        self.assertEqual(shift_formula_string("=SUM(F10:F25)", insert_row=15, delta=10), "=SUM(F10:F35)")
        # Spanning column range: F..L (cols 6..12), insert at 8 with delta 3 -> F..O (cols 6..15)
        self.assertEqual(shift_formula_string("=SUM(F10:L10)", insert_col=8, delta_cols=3), "=SUM(F10:O10)")
        # Downstream column shifting: M..N (13..14) -> P..Q (16..17), Z (26) -> AC (29)
        self.assertEqual(shift_formula_string("=SUM(M10:N10) + Z5", insert_col=8, delta_cols=3), "=SUM(P10:Q10) + AC5")
        # Cross-sheet references with single-quoted sheet prefix
        self.assertEqual(
            shift_formula_string("='NAV-SVC'!$F$34:$H$34", insert_row=20, delta_rows=5),
            "='NAV-SVC'!$F$39:$H$39"
        )

    def test_horizontal_column_expansion(self):
        """Verify expand_table_columns clones column width, header/cell formatting, and shifts formulas."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Matrix"

        # Setup base columns
        for c in range(1, 10):
            ws.cell(row=1, column=c, value=f"Col_{c}")
        # Col 6 (Col F) has prototype width and styling
        ws.column_dimensions["F"].width = 22.0
        ws.cell(row=1, column=6).font = Font(name="Segoe UI", bold=True, size=11)
        ws.cell(row=2, column=6, value="O")
        ws.cell(row=2, column=6).alignment = Alignment(horizontal="center")
        # Summary formula summing across cols F..I (cols 6..9)
        ws.cell(row=5, column=1, value="=SUM(F2:I2)")

        engine = UniversalExcelEngine(wb)
        # Expand 3 columns at column 7 (insert between Col F and G)
        new_cols = engine.expand_table_columns("Matrix", start_col=7, count=3, prototype_col=6)
        self.assertEqual(new_cols, [7, 8, 9])

        # Verify cloned width on newly inserted columns
        self.assertEqual(ws.column_dimensions["G"].width, 22.0)
        self.assertEqual(ws.column_dimensions["H"].width, 22.0)
        self.assertEqual(ws.column_dimensions["I"].width, 22.0)

        # Verify formula expanded horizontally: F2:I2 -> F2:L2 (original 6..9 expanded by 3 to 6..12)
        new_formula = ws.cell(row=5, column=1).value
        self.assertEqual(new_formula, "=SUM(F2:L2)")

    def test_chart_reanchoring_and_series_relink(self):
        """Verify automated chart re-anchoring and series formula updates (Invariant E12)."""
        from openpyxl.chart import PieChart, Reference

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ChartSheet"

        ws["A1"] = "Status"
        ws["B1"] = "Count"
        ws["A2"] = "Pass"
        ws["B2"] = 25
        ws["A3"] = "Fail"
        ws["B3"] = 5
        ws["B4"] = "=SUM(B2:B3)"

        chart = PieChart()
        chart.title = "Test Distribution"
        data_ref = Reference(ws, min_col=2, min_row=1, max_row=3)
        chart.add_data(data_ref, titles_from_data=True)
        # Anchor at D10
        chart.anchor = "D10"
        ws.add_chart(chart)

        engine = UniversalExcelEngine(wb)
        # Insert 6 rows at row 2
        engine.expand_table_rows("ChartSheet", start_row=2, count=6)

        # Verify chart anchor shifted down by 6 rows: D10 -> D16
        self.assertEqual(chart.anchor, "D16")
        # Verify chart series formula shifted
        series_formula = chart.series[0].val.numRef.f
        self.assertEqual(series_formula, "'ChartSheet'!$B$2:$B$9")

    def test_headless_openpyxl_reader(self):
        """Verify headless openpyxl reader produces structured JSON snapshot without Microsoft Excel."""
        from xlsx_reader import read_xlsx

        mock_path = os.path.join(self.output_dir, "mock_template.xlsx")
        self.assertTrue(os.path.isfile(mock_path))

        snapshot = read_xlsx(mock_path, engine="openpyxl")
        self.assertIn("source_file", snapshot)
        self.assertIn("sheets", snapshot)
        self.assertGreater(len(snapshot["sheets"]), 0)

        sheet0 = snapshot["sheets"][0]
        self.assertEqual(sheet0["sheet_name"], "Test Cases")
        self.assertIn("cells", sheet0)

        # Find cell A1
        a1_cell = next((c for c in sheet0["cells"] if c["address"] == "A1"), None)
        self.assertIsNotNone(a1_cell)
        self.assertIn("PROJECT:", a1_cell["value"])
        self.assertTrue(a1_cell["format"]["bold"])

    def test_12_quality_gates_validator(self):
        """Verify 12 Quality Gates validation engine (Invariant E6)."""
        from xlsx_validator import validate_excel_sheet

        asset_path = os.path.join(BASE_DIR, "excel_assets", "Report5_Unit Test_Enhanced_UX_v2.xlsx")
        if os.path.isfile(asset_path):
            # Test that validator detects known differences between Function 1 and Example
            res = validate_excel_sheet(asset_path, target_sheet="Function 1", ref_sheet="Example", verbose=False)
            self.assertEqual(res.target_sheet, "Function 1")
            self.assertEqual(res.ref_sheet, "Example")
            self.assertEqual(res.gates_evaluated, 12)
            self.assertFalse(res.passed)
            self.assertGreater(len(res.diffs), 0)
            self.assertGreater(res.gates_failed, 0)
        else:
            mock_path = os.path.join(self.output_dir, "mock_template.xlsx")
            res = validate_excel_sheet(mock_path, target_sheet="Test Cases", ref_sheet="Test Cases", verbose=False)
            self.assertEqual(res.gates_evaluated, 12)


if __name__ == "__main__":
    unittest.main()
