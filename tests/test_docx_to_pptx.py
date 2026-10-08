"""
Unit test suite for Word (.docx) to PowerPoint (.pptx) Cross-Format Pipeline.
Tests docx_to_pptx.py, smart chunking, table pagination, diagram extraction, and thesis_blue theme.
"""

from __future__ import annotations

import json
import os
import tempfile
import unittest

from PIL import Image
from docx import Document
from docx.shared import Inches, Pt
from pptx import Presentation

from docx_to_pptx import extract_docx_structure, synthesize_slide_spec, convert_docx_to_pptx
from pptx_writer import THEMES


class TestDocxToPptxPipeline(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dpath = self.temp_dir.name

        # Create a sample test image
        self.sample_img_path = os.path.join(self.dpath, "sample_diagram.png")
        im = Image.new("RGB", (800, 400), color=(0, 81, 226))
        im.save(self.sample_img_path)

        # Build a rich sample Word document
        self.sample_docx = os.path.join(self.dpath, "sample_report.docx")
        doc = Document()
        doc.core_properties.title = "Báo Cáo Nghiên Cứu Hệ Thống IoT"
        doc.core_properties.author = "TS. Vũ Nguyễn"

        # Chapter 1
        p_ch1 = doc.add_paragraph("Chương 1. Giới thiệu tổng quan hệ thống", style="Heading 1")
        p_desc1 = doc.add_paragraph(
            "Hệ thống giám sát năng lượng thông minh được xây dựng nhằm mục đích tối ưu hóa hiệu suất tiêu thụ điện năng. "
            "Giải pháp áp dụng các thuật toán học máy kết hợp mạng cảm biến không dây để thu thập và phân tích dữ liệu theo thời gian thực. "
            "Dữ liệu được mã hóa chuẩn công nghiệp và lưu trữ tập trung tại cơ sở dữ liệu phân tán an toàn."
        )
        doc.add_paragraph("• Khả năng chịu lỗi và tính sẵn sàng cao (High Availability).")
        doc.add_paragraph("• Giao thức truyền thông không dây băng thông thấp tiết kiệm năng lượng.")
        doc.add_paragraph("• Tích hợp hệ thống cảnh báo tức thời qua ứng dụng di động.")

        # Subsection 1.1 with Table
        doc.add_paragraph("1.1. Bảng thông số kỹ thuật cảm biến", style="Heading 2")
        tbl = doc.add_table(rows=4, cols=3)
        headers = ["STT", "Tên Thông Số", "Giá Trị"]
        for c_idx, h in enumerate(headers):
            tbl.cell(0, c_idx).text = h
        data = [
            ["01", "Điện áp hoạt động", "3.3V DC"],
            ["02", "Dòng tiêu thụ", "15mA"],
            ["03", "Sai số đo lường", "< 0.5%"],
        ]
        for r_idx, row_vals in enumerate(data, 1):
            for c_idx, val in enumerate(row_vals):
                tbl.cell(r_idx, c_idx).text = val

        # Subsection 1.2 with Diagram & Caption
        doc.add_paragraph("1.2. Sơ đồ khối kiến trúc phần cứng", style="Heading 2")
        p_img = doc.add_paragraph()
        p_img.add_run().add_picture(self.sample_img_path, width=Inches(4.5))
        doc.add_paragraph("Hình 1.1: Sơ đồ kiến trúc vi điều khiển và khối thu phát không dây.")

        # Chapter 2 with Large Table (>7 rows for pagination testing)
        doc.add_paragraph("Chương 2. Kết quả kiểm thử và danh mục linh kiện", style="Heading 1")
        doc.add_paragraph("2.1. Danh mục linh kiện phần cứng chi tiết", style="Heading 2")

        big_tbl = doc.add_table(rows=11, cols=3)
        big_headers = ["Mã LK", "Tên Linh Kiện", "Số Lượng"]
        for c_idx, h in enumerate(big_headers):
            big_tbl.cell(0, c_idx).text = h
        for i in range(1, 11):
            big_tbl.cell(i, 0).text = f"IC-00{i}"
            big_tbl.cell(i, 1).text = f"Module Linh Kiện {i}"
            big_tbl.cell(i, 2).text = str(i * 2)

        doc.save(self.sample_docx)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_theme_registration(self):
        """Validates thesis_blue theme exists with Tahoma font, Pure Black title (PPTX_INV_10), and Royal Blue accent."""
        self.assertIn("thesis_blue", THEMES)
        theme = THEMES["thesis_blue"]
        self.assertEqual(theme["font_heading"], "Tahoma")
        self.assertEqual(theme["title_color"], (0, 0, 0))
        self.assertEqual(theme["accent_color"], (0, 81, 226))

    def test_extract_docx_structure(self):
        """Validates AST extraction of chapters, subsections, tables, and images."""
        tree = extract_docx_structure(self.sample_docx, asset_dir=os.path.join(self.dpath, "extracted_assets"))
        self.assertEqual(tree["title"], "Báo Cáo Nghiên Cứu Hệ Thống IoT")
        self.assertEqual(tree["author"], "TS. Vũ Nguyễn")
        self.assertGreaterEqual(len(tree["chapters"]), 2)

        # Check chapter 1
        ch1 = tree["chapters"][0]
        self.assertIn("Chương 1", ch1["title"])
        self.assertGreaterEqual(len(ch1["subsections"]), 2)

    def test_synthesize_slide_spec_smart_pagination(self):
        """Validates slide spec generation, 4-5 bullet limits, and table pagination (>7 rows)."""
        tree = extract_docx_structure(self.sample_docx, asset_dir=os.path.join(self.dpath, "extracted_assets"))
        spec = synthesize_slide_spec(tree, theme="thesis_blue", max_bullets_per_slide=4, max_table_rows_per_slide=7)

        self.assertEqual(spec["theme"], "thesis_blue")
        slides = spec["slides"]
        self.assertGreaterEqual(len(slides), 6)

        # Slide 1: Title
        self.assertEqual(slides[0]["type"], "title")
        self.assertEqual(slides[0]["title"], "Báo Cáo Nghiên Cứu Hệ Thống IoT")

        # Slide 2: Agenda Grid Cards
        self.assertEqual(slides[1]["type"], "grid_cards")
        self.assertIn("AGENDA", slides[1]["title"].upper())

        # Check Table Pagination: Big table with 10 rows should split into 2 table slides
        table_slides = [s for s in slides if s.get("type") == "table"]
        self.assertGreaterEqual(len(table_slides), 3)  # 1 small table + 2 pages of big table

        paginated = [s for s in table_slides if "Bảng trang" in s.get("title", "")]
        self.assertEqual(len(paginated), 2)
        self.assertEqual(len(paginated[0]["rows"]), 7)
        self.assertEqual(len(paginated[1]["rows"]), 3)

    def test_end_to_end_conversion(self):
        """Validates end-to-end execution: docx ➔ slide_spec.json ➔ finished .pptx."""
        out_pptx = os.path.join(self.dpath, "final_presentation.pptx")
        spec_out = os.path.join(self.dpath, "final_spec.json")

        final_pptx, final_spec = convert_docx_to_pptx(
            docx_path=self.sample_docx,
            output_pptx=out_pptx,
            spec_out=spec_out,
            theme="thesis_blue"
        )

        self.assertTrue(os.path.isfile(final_pptx))
        self.assertTrue(os.path.isfile(final_spec))

        # Inspect generated PPTX presentation with python-pptx
        prs = Presentation(final_pptx)
        self.assertGreaterEqual(len(prs.slides), 7)
        # Check widescreen 16:9
        self.assertAlmostEqual(prs.slide_width.inches, 13.333, places=2)
        self.assertAlmostEqual(prs.slide_height.inches, 7.5, places=2)


if __name__ == "__main__":
    unittest.main()
