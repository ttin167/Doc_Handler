# -*- coding: utf-8 -*-
"""
tests/test_core_algorithms.py
Automated Unit & Integration Test Suite for:
1. Heuristic Table Column Auto-Sizing (Square Root Damped Weighting in Word).
2. Auto-Pagination & Dynamic Font Downsizing in PowerPoint.
"""

import os
import tempfile
import unittest
import docx
from docx.shared import Inches

from smart_post_processor import (
    calculate_heuristic_column_widths,
    apply_heuristic_table_widths,
)
from markdown_converter import markdown_to_docx
from pptx_writer import (
    paginate_slide_specs,
    write_pptx_from_spec,
    add_content_slide,
    init_presentation,
)


class TestCoreAlgorithms(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    # -------------------------------------------------------------------------
    # Algorithm 1: Heuristic Table Column Auto-Sizing Tests
    # -------------------------------------------------------------------------

    def test_calculate_heuristic_column_widths_sum_and_bounds(self):
        """Verifies exact sum constraint and min/max clamps in column weighting."""
        table_data = [
            ["ID", "Mã", "Tên", "Mô tả nghiệp vụ chi tiết"],
            ["1", "PRD-01", "Bánh mì", "Bánh mì kẹp thịt tươi ngon giàu dinh dưỡng cho bữa sáng năng động của học sinh sinh viên"],
            ["2", "PRD-02", "Sữa tươi", "Sữa tươi thanh trùng 100% nguyên chất tiệt trùng từ trang trại bò sữa tự nhiên"],
            ["3", "PRD-03", "Trà đào", "Trà đào thanh mát với đào miếng giòn rụm kết hợp sả tắc tạo hương vị độc đáo tươi mát"]
        ]
        total_twips = 9360  # ~6.5 inches printable width

        widths = calculate_heuristic_column_widths(table_data, total_width_twips=total_twips)

        self.assertEqual(len(widths), 4)
        # 1. Exact sum invariant
        self.assertEqual(sum(widths), total_twips)

        # 2. Asymmetric distribution: Description column (idx 3) must be widest
        self.assertTrue(widths[0] < widths[1] <= widths[2] < widths[3])

        # 3. Min clamp: ID column (idx 0) must not collapse below 8% (~748 twips)
        self.assertGreaterEqual(widths[0], int(total_twips * 0.08))

        # 4. Max clamp: Description column must not exceed 55%
        self.assertLessEqual(widths[3], int(total_twips * 0.55))

    def test_markdown_to_docx_applies_heuristic_column_widths(self):
        """Verifies end-to-end markdown table conversion produces balanced non-equal columns."""
        md_text = """# Báo Cáo Sản Phẩm

| STT | Mã SP | Tên Mặt Hàng | Diễn Giải Chi Tiết Chức Năng |
|:---:|:---:|:---|:---|
| 1 | A01 | Cam sành | Cam sành loại 1 mọng nước ngọt thanh xuất khẩu |
| 2 | A02 | Bưởi da xanh | Bưởi da xanh ruột hồng ngọt dịu không hạt chuẩn VietGAP |
"""
        md_path = os.path.join(self.temp_dir, "table_test.md")
        docx_path = os.path.join(self.temp_dir, "table_test.docx")

        with open(md_path, "w", encoding="utf-8") as f:
            f.write(md_text)

        ok, out_path = markdown_to_docx(md_path, docx_path)
        self.assertTrue(ok)
        self.assertTrue(os.path.isfile(out_path))

        doc = docx.Document(out_path)
        self.assertEqual(len(doc.tables), 1)
        tbl = doc.tables[0]

        # Verify autofit is False for strict geometry
        self.assertFalse(tbl.autofit)

        # Verify cells have widths applied and description is significantly wider than STT
        stt_w = tbl.rows[0].cells[0].width
        desc_w = tbl.rows[0].cells[3].width
        self.assertGreater(desc_w, stt_w * 1.5)

    # -------------------------------------------------------------------------
    # Algorithm 2: PowerPoint Auto-Pagination & Dynamic Font Downsizing Tests
    # -------------------------------------------------------------------------

    def test_pptx_auto_pagination_splits_dense_bullets(self):
        """Verifies that slides with > 7 bullets automatically split with the SAME title (no suffix)."""
        slides_spec = [
            {
                "type": "content",
                "title": "KIẾN TRÚC HỆ THỐNG",
                "bullets": [
                    f"Gạch đầu dòng phân tích thành phần hệ thống số {i}: Đảm bảo độ trễ thấp và tính sẵn sàng cao."
                    for i in range(1, 13)  # 12 bullets
                ]
            }
        ]

        paginated = paginate_slide_specs(slides_spec, max_bullets_threshold=7)

        # Must split into 2 slides (since 12 > 7)
        self.assertEqual(len(paginated), 2)

        # Both slides MUST have the EXACT SAME title (no '(Tiếp theo)' or '(1/2)' per user requirement)
        self.assertEqual(paginated[0]["title"], "KIẾN TRÚC HỆ THỐNG")
        self.assertEqual(paginated[1]["title"], "KIẾN TRÚC HỆ THỐNG")

        # Total bullets preserved exactly
        self.assertEqual(len(paginated[0]["bullets"]) + len(paginated[1]["bullets"]), 12)
        self.assertLessEqual(len(paginated[0]["bullets"]), 6)
        self.assertLessEqual(len(paginated[1]["bullets"]), 6)

    def test_pptx_auto_pagination_end_to_end_build(self):
        """Verifies write_pptx_from_spec compiles auto-paginated deck without crashing."""
        out_pptx = os.path.join(self.temp_dir, "auto_paginated_deck.pptx")
        spec = {
            "theme": "thesis_blue",
            "slides": [
                {
                    "type": "title",
                    "title": "BÁO CÁO NGHIÊN CỨU",
                    "subtitle": "Thử nghiệm thuật toán Auto-Pagination"
                },
                {
                    "type": "content",
                    "title": "DANH MỤC THỬ NGHIỆM ĐỒNG BỘ",
                    "bullets": [
                        f"Mục tiêu kiểm thử chất lượng hệ thống {i}: Xác nhận tính toàn vẹn của mô-đun."
                        for i in range(1, 11)  # 10 bullets
                    ]
                }
            ]
        }

        final_path = write_pptx_from_spec(spec, out_pptx)
        self.assertTrue(os.path.isfile(final_path))

        # Check presentation slide count: 1 Title + 2 Content (split from 1) = 3 Slides
        from pptx import Presentation
        prs = Presentation(final_path)
        self.assertEqual(len(prs.slides), 3)


if __name__ == "__main__":
    unittest.main()
